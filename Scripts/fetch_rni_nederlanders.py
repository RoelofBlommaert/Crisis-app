"""
Pull real, publicly-published counts of Dutch nationals registered as
living abroad, per tracked country.

Data source: CBS "Nederlanders in het buitenland" maatwerktabel
(https://www.cbs.nl/nl-nl/maatwerk/2025/40/nederlanders-in-het-buitenland-1-juli-2025),
a custom CBS publication derived from the RNI (Register
Niet-Ingezetenen) via the BRP. This is NOT raw internal RNI/BRP data --
it is CBS's own aggregated, already-de-identified, publicly downloadable
release (an .xlsx, no login, no key). Confirmed by downloading and
inspecting the actual file before writing this script, not assumed:
all 8 tracked countries appear as individually-named rows (CBS presents
167 countries individually, not bucketed into regions the way its old
tourism table was -- see the now-removed fetch_cbs_travel.py).

IMPORTANT CBS caveat (from the file's own "Toelichting" sheet, carried
through into every row of this script's output as `cbs_note`): these
counts are "vrijwel altijd een onderschatting" (almost always an
underestimate) of the real number of Dutch nationals abroad, because
RNI registration is not mandatory for non-residents. CBS's own example:
Canada shows ~19k RNI-registered Dutch nationals here, vs. ~89k
Netherlands-born residents in Statistics Canada's own census -- a 4-5x
gap. Treat this as a floor/registered-count, not a total population
estimate.

CBS does not publish a stable "latest" URL for this maatwerktabel --
each edition gets a new URL under /maatwerk/<year>/<week>/. This script
pins CBS_XLSX_URL to the most recent edition found (1 July 2025);
update it by hand when a newer edition is published (CBS has released
these roughly twice a year -- confirmed editions exist for July 2022,
January 2025, and July 2025).

Output:
  Data/RNI/rni_nederlanders_<date>.csv
"""

import csv
import io
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
import requests

from countries import COUNTRIES

CBS_XLSX_URL = "https://www.cbs.nl/-/media/_excel/2025/40/nederlanders-in-het-buitenland-20250701.xlsx"
SOURCE_DATE = "2025-07-01"  # the peildatum printed on the CBS "Voorblad" sheet for this edition

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = REPO_ROOT / "Data" / "RNI"

CBS_NOTE = (
    "CBS: cijfers zijn vrijwel altijd een onderschatting van het werkelijke "
    "aantal Nederlanders in het buitenland, omdat RNI-registratie voor "
    "niet-ingezetenen niet verplicht is (bv. Canada: 19k RNI-geregistreerden "
    "vs. bijna 89k in Nederland geboren inwoners volgens Statistics Canada)."
)


def fetch_workbook() -> openpyxl.Workbook:
    resp = requests.get(CBS_XLSX_URL, timeout=30)
    resp.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)


def parse_table1(wb: openpyxl.Workbook) -> dict:
    """Row shape (Tabel 1): continent, country, born_in_nl, born_in_country,
    born_elsewhere, born_country_of_unknown_residence, total."""
    ws = wb["Tabel 1"]
    by_country = {}
    for row in ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=True):
        country = row[1]
        if not country:
            continue
        by_country[country] = {
            "born_in_nl": row[2] or 0,
            "born_in_country": row[3] or 0,
            "born_elsewhere": row[4] or 0,
            "total": row[6] or 0,
        }
    return by_country


def main() -> None:
    print(f"Downloading CBS RNI table from {CBS_XLSX_URL} ...")
    wb = fetch_workbook()
    by_country = parse_table1(wb)
    print(f"  parsed {len(by_country)} countries from Tabel 1")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / f"rni_nederlanders_{date.today().isoformat()}.csv"

    rows = []
    for c in COUNTRIES:
        cbs_name = c["cbs_rni_name"]
        stats = by_country.get(cbs_name)
        if stats is None:
            print(f"  ! no CBS row found for {c['name']} (looked for '{cbs_name}')")
            continue
        rows.append(
            {
                "country": c["name"],
                "cbs_rni_name": cbs_name,
                "rni_registered": stats["total"],
                "born_in_nl": stats["born_in_nl"],
                "born_in_country": stats["born_in_country"],
                "born_elsewhere": stats["born_elsewhere"],
                "source_date": SOURCE_DATE,
                "fetched_at": datetime.now(timezone.utc).isoformat(),
                "cbs_note": CBS_NOTE,
            }
        )
        print(f"  {c['name']:10s} ({cbs_name}): {stats['total']} RNI-registered")

    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "country",
                "cbs_rni_name",
                "rni_registered",
                "born_in_nl",
                "born_in_country",
                "born_elsewhere",
                "source_date",
                "fetched_at",
                "cbs_note",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    print(f"\nSaved {len(rows)} countries to {out_path}")


if __name__ == "__main__":
    main()
